import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import ActualRecord, AuditEvent, Milestone, PlanItem, PlanVersion, Project
from app.seed import load_sample
from conftest import REVIEWER_ID, as_reviewer


def payload(code="P-TEST"):
    return {
        "project_code": code,
        "project_name": "验证项目",
        "project_type": "船舶建造",
        "owner_id": "00000000-0000-0000-0000-000000000002",
        "department_id": "00000000-0000-0000-0000-000000000001",
    }


def test_racing_editors_cannot_overwrite_another_committed_project_change(client, factory):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    p, _ = create(client, "P-RACE")
    barrier = Barrier(2)

    def edit(name):
        barrier.wait(timeout=10)
        return name, client.patch(
            f"/api/v1/projects/{p['id']}",
            headers={"If-Match": str(p["row_version"])},
            json={**payload("P-RACE"), "project_name": name},
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(edit, ["并发编制A", "并发编制B"]))
    assert sorted(r.status_code for _, r in responses) == [200, 409]
    winner = next(name for name, r in responses if r.status_code == 200)
    with factory() as db:
        assert db.get(Project, p["id"]).project_name == winner
        assert (
            db.scalar(
                select(func.count())
                .select_from(AuditEvent)
                .where(AuditEvent.aggregate_id == p["id"], AuditEvent.action == "update")
            )
            == 1
        ), "只有获胜的事务留下修改审计；失败请求不能部分提交"


def create(client, code="P-TEST"):
    response = client.post("/api/v1/projects", json=payload(code))
    assert response.status_code == 201, response.text
    p = response.json()["data"]
    assert (
        client.put(
            f"/api/v1/projects/{p['id']}/members",
            json={"user_id": REVIEWER_ID, "role_code": "reviewer"},
        ).status_code
        == 200
    )
    plan = client.get(f"/api/v1/projects/{p['id']}/plan-versions").json()["data"][0]
    return p, plan


def node_payload(plan_id, code="N-001", planned="2026-10-01"):
    return {
        "plan_version_id": plan_id,
        "milestone_code": code,
        "milestone_name": "图纸评审",
        "node_type_code": "设计资料",
        "planned_date": planned,
        "department_id": "00000000-0000-0000-0000-000000000001",
        "change_reason": "初始录入",
    }


def excel(rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.append(headers or ["项目编号", "项目名称"])
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def preview(client, rows):
    return client.post(
        "/api/v1/imports/preview",
        files={"file": ("test.xlsx", excel(rows))},
        data={
            "sheet_name": "data",
            "mapping_json": '{"project_code":"项目编号","project_name":"项目名称"}',
        },
    )


def commit(client, batch, key="test-key", **overrides):
    return client.post(
        f"/api/v1/imports/{batch['id']}/commit",
        headers={"Idempotency-Key": key},
        json={"preview_hash": batch["preview_hash"], "include_valid_only": False, **overrides},
    )


def test_project_crud_unique_lock_and_persistence(client, factory):
    """A second editor must never silently overwrite a colleague's project update."""
    p, _ = create(client)
    assert client.post("/api/v1/projects", json=payload()).status_code == 409
    changed = {**payload(), "project_name": "新名称", "project_status": "active"}
    response = client.patch(f"/api/v1/projects/{p['id']}", headers={"If-Match": "1"}, json=changed)
    assert response.status_code == 200
    assert (
        client.patch(
            f"/api/v1/projects/{p['id']}", headers={"If-Match": "1"}, json=payload()
        ).status_code
        == 409
    )
    result = client.get(
        "/api/v1/projects?status=active&type=船舶建造&keyword=张工&page_size=1"
    ).json()["data"]
    assert result["total"] == 1 and result["items"][0]["project_name"] == "新名称"
    # New independent session reads the committed write, rather than the request identity map.
    with factory() as db:
        assert db.get(Project, p["id"]).project_name == "新名称"
        assert db.scalar(select(func.count()).select_from(AuditEvent)) >= 2
    archived = client.delete(f"/api/v1/projects/{p['id']}", headers={"If-Match": "2"})
    assert archived.json()["data"]["project_status"] == "archived"
    assert client.get(f"/api/v1/projects/{p['id']}").status_code == 200


def test_database_prevents_contradictory_actual_facts(client, factory):
    """A future integration bypassing API validation cannot date an incomplete node."""
    from datetime import date

    from sqlalchemy.exc import IntegrityError

    p, plan = create(client)
    node = client.post(
        f"/api/v1/projects/{p['id']}/milestones", json=node_payload(plan["id"])
    ).json()["data"]
    with pytest.raises(IntegrityError), factory.begin() as db:
        db.add(
            ActualRecord(
                milestone_id=node["id"],
                actual_date=date(2026, 10, 1),
                completion_flag=False,
                evidence_ref="contradictory fact",
            )
        )


def test_validation_dates_and_foreign_keys(client):
    assert (
        client.post(
            "/api/v1/projects",
            json={**payload(), "planned_start": "2026-10-02", "planned_end": "2026-10-01"},
        ).status_code
        == 422
    )
    assert (
        client.post("/api/v1/projects", json={**payload(), "department_id": "unknown"}).status_code
        == 422
    )
    assert (
        client.post("/api/v1/projects", json={**payload(), "project_name": "  "}).status_code == 422
    )


def test_project_pagination_has_no_missing_or_repeated_rows(client):
    """Pagination must preserve every filtered project exactly once."""
    for code in ("P-1", "P-2", "P-3", "P-4"):
        create(client, code)
    first = client.get(
        "/api/v1/projects?page_size=2&page=1&status=not_started&type=船舶建造"
    ).json()["data"]
    second = client.get(
        "/api/v1/projects?page_size=2&page=2&status=not_started&type=船舶建造"
    ).json()["data"]
    assert first["total"] == second["total"] == 4
    assert len({p["id"] for p in first["items"] + second["items"]}) == 4


def test_plan_snapshots_are_immutable_and_full_copy(client, factory):
    """A schedule change must preserve the old date and unchanged nodes for traceability."""
    p, plan = create(client)
    for code in ("N-001", "N-002"):
        assert (
            client.post(
                f"/api/v1/projects/{p['id']}/milestones", json=node_payload(plan["id"], code)
            ).status_code
            == 201
        )
    with factory.begin() as db:
        db.get(PlanVersion, plan["id"]).source_ref = "原始计划资料"
    plan = client.get(f"/api/v1/projects/{p['id']}/plan-versions").json()["data"][0]
    with as_reviewer(client):
        assert (
            client.post(
                f"/api/v1/plan-versions/{plan['id']}/activate",
                headers={"If-Match": str(plan["row_version"])},
                json={"approval_ref": "演示"},
            ).status_code
            == 200
        )
    nodes = client.get(f"/api/v1/projects/{p['id']}/milestones").json()["data"]["items"]
    approved = client.get(f"/api/v1/projects/{p['id']}/plan-versions").json()["data"][0]
    assert approved["source_ref"] == "原始计划资料" and approved["approval_ref"] == "演示"
    old_node = nodes[0]
    assert (
        client.patch(
            f"/api/v1/milestones/{old_node['id']}/plan-item",
            headers={"If-Match": str(old_node["row_version"])},
            json=node_payload(plan["id"], planned="2026-10-05"),
        ).status_code
        == 409
    )
    p = client.get(f"/api/v1/projects/{p['id']}").json()["data"]
    draft = client.post(
        f"/api/v1/projects/{p['id']}/plan-versions",
        headers={"If-Match": str(p["row_version"])},
        json={"change_reason": "设计变更"},
    ).json()["data"]
    assert draft["version_label"] == "V2"
    updated = client.patch(
        f"/api/v1/milestones/{old_node['id']}/plan-item",
        headers={"If-Match": str(old_node["row_version"])},
        json=node_payload(draft["id"], planned="2026-10-05"),
    )
    assert updated.status_code == 200, updated.text
    with factory() as db:
        items = db.scalars(select(PlanItem).where(PlanItem.milestone_id == old_node["id"])).all()
        assert sorted(str(i.planned_date) for i in items) == ["2026-10-01", "2026-10-05"]
        assert (
            db.scalar(
                select(func.count())
                .select_from(PlanItem)
                .where(PlanItem.plan_version_id == draft["id"])
            )
            == 2
        )
    assert (
        client.patch(
            f"/api/v1/milestones/{old_node['id']}/plan-item",
            headers={"If-Match": str(old_node["row_version"])},
            json=node_payload(draft["id"]),
        ).status_code
        == 409
    )


def test_actual_submission_is_not_completion_and_keeps_prior_verified_fact(client):
    """New feedback cannot replace a verified fact until explicit independent verification."""
    p, plan = create(client)
    node = client.post(
        f"/api/v1/projects/{p['id']}/milestones", json=node_payload(plan["id"])
    ).json()["data"]
    path = f"/api/v1/milestones/{node['id']}/actuals"
    first = client.post(
        path,
        json={"actual_date": "2026-10-01", "completion_flag": True, "evidence_ref": "演示记录"},
    ).json()["data"]
    assert first["verification_status"] == "submitted" and not first["is_current"]
    detail = client.get(f"/api/v1/milestones/{node['id']}").json()["data"]
    assert detail["current_actual"] is None
    with as_reviewer(client):
        result = client.post(
            f"/api/v1/milestone-actuals/{first['id']}/verify",
            json={"decision": "verified", "comment": "演示核实"},
        )
    assert result.status_code == 200
    second = client.post(
        path, json={"completion_flag": None, "evidence_ref": "新信息待核对"}
    ).json()["data"]
    detail = client.get(f"/api/v1/milestones/{node['id']}").json()["data"]
    assert detail["current_actual"]["id"] == first["id"]
    assert detail["verification_status"] == "submitted"
    with as_reviewer(client):
        assert (
            client.post(
                f"/api/v1/milestone-actuals/{second['id']}/verify",
                json={"decision": "returned", "comment": "补证据"},
            ).status_code
            == 200
        )
    assert (
        client.post(path, json={"completion_flag": True, "evidence_ref": "缺日期"}).status_code
        == 422
    )


def test_import_preview_is_readonly_and_commit_idempotent(client, factory):
    """Preview and double-clicking confirm must never create duplicate business rows."""
    batch = preview(client, [["P-A", "A"], ["P-B", "B"]]).json()["data"]
    assert batch["valid_rows"] == 2
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(Project)) == 0
    assert commit(client, batch).status_code == 200
    assert commit(client, batch).json()["data"]["result_json"]["inserted"] == 2
    assert commit(client, batch, include_valid_only=True).status_code == 409
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(Project)) == 2


def test_concurrent_confirm_never_duplicates_and_postgres_replays_result(client, factory):
    """Concurrent confirm clicks must serialize and leave exactly one business record."""
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    batch = preview(client, [["P-CONCURRENT", "并发确认"]]).json()["data"]
    barrier = Barrier(2)

    def confirm():
        barrier.wait(timeout=10)
        return commit(client, batch, key="concurrent-key")

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: confirm(), range(2)))
    assert all(r.status_code in (200, 409) for r in responses)
    if factory.kw["bind"].dialect.name == "postgresql":
        assert [r.status_code for r in responses] == [200, 200]
    with factory() as db:
        assert (
            db.scalar(
                select(func.count())
                .select_from(Project)
                .where(Project.project_code == "P-CONCURRENT")
            )
            == 1
        )


def test_import_errors_atomic_and_downloadable(client):
    create(client, "P-A")
    result = preview(
        client, [["P-A", "已有"], ["P-B", "合法"], ["P-B", "同表重复"], [None, "缺编号"]]
    )
    assert result.status_code == 201
    batch = result.json()["data"]
    assert batch["valid_rows"] == 1 and batch["error_rows"] == 3
    assert commit(client, batch).status_code == 409
    assert client.get("/api/v1/projects").json()["data"]["total"] == 1
    assert "DUPLICATE" in client.get(f"/api/v1/imports/{batch['id']}/errors").text


def test_import_rechecks_conflicts_after_preview(client):
    batch = preview(client, [["P-RACE", "导入前"]]).json()["data"]
    create(client, "P-RACE")
    assert commit(client, batch).status_code == 409
    assert client.get("/api/v1/projects").json()["data"]["total"] == 1


@pytest.mark.parametrize("name,content", [("test.csv", b"a,b"), ("test.xlsx", b"not-excel")])
def test_invalid_files_are_rejected(client, name, content):
    assert (
        client.post("/api/v1/imports/preview", files={"file": (name, content)}).status_code == 400
    )


def test_formula_input_rejected(client):
    batch = preview(client, [["P-F", "=1+1"]]).json()["data"]
    assert batch["error_rows"] == 1


def test_node_import_to_draft_and_cross_project_rejected(client):
    p, plan = create(client)
    content = excel(
        [[p["project_code"], "N-X", "导入节点", "设计资料", "2026-10-20"]],
        ["项目编号", "节点编号", "节点名称", "节点类型", "计划完成日期"],
    )
    mapping = '{"project_code":"项目编号","milestone_code":"节点编号","milestone_name":"节点名称","node_type_code":"节点类型","planned_date":"计划完成日期"}'
    result = client.post(
        "/api/v1/imports/preview",
        files={"file": ("nodes.xlsx", content)},
        data={
            "target": "milestones",
            "mapping_json": mapping,
            "project_id": p["id"],
            "plan_version_id": plan["id"],
        },
    )
    assert result.status_code == 201, result.text
    batch = result.json()["data"]
    assert batch["valid_rows"] == 1
    assert commit(client, batch).status_code == 200
    other, other_plan = create(client, "P-OTHER")
    invalid = client.post(
        f"/api/v1/projects/{p['id']}/milestones", json=node_payload(other_plan["id"])
    )
    assert invalid.status_code == 422


def test_sample_workbook_mapping_and_no_invented_facts(factory):
    """The supplied 67 source nodes, dates and verification flags must survive loading intact."""
    path = Path(__file__).resolve().parents[2] / "模拟业务测试数据 V0.1.xlsx"
    with factory.begin() as db:
        assert load_sample(db, path) == 3
        assert load_sample(db, path) == 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        with factory() as db:
            for row in list(workbook["02_节点计划"].values)[6:]:
                p = db.scalar(select(Project).where(Project.project_code == row[0]))
                n = db.scalar(
                    select(Milestone).where(
                        Milestone.project_id == p.id, Milestone.milestone_code == row[1]
                    )
                )
                item = db.scalar(select(PlanItem).where(PlanItem.milestone_id == n.id))
                assert item.planned_date == row[5].date()
                actual = db.scalar(
                    select(ActualRecord).where(
                        ActualRecord.milestone_id == n.id, ActualRecord.is_current.is_(True)
                    )
                )
                if row[7] == "已核实":
                    assert actual.actual_date == row[6].date()
                else:
                    assert actual is None
            assert db.scalar(select(func.count()).select_from(Milestone)) == 67
            p = db.scalar(select(Project).where(Project.project_code == "P-H028"))
            assert db.get(PlanVersion, p.current_plan_version_id).version_label == "V2"
    finally:
        workbook.close()
