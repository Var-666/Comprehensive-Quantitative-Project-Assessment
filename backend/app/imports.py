"""Deterministic Excel mapping, preview and transactional commit. No AI in the write path."""

import hashlib
import io
import json
from datetime import date, datetime
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from types import SimpleNamespace

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import select

from .config import settings
from .models import (
    AuditEvent,
    Department,
    ImportBatch,
    ImportRow,
    Milestone,
    PlanItem,
    PlanVersion,
    Project,
    User,
    now,
)
from .schemas import MilestoneInput, ProjectInput
from . import action_matching, node_configuration
from . import workflow as wf
from .workflow_models import ActionTemplate
from .services import audit, create_milestone, create_project, fail, get, record

PROJECT_MAPPING = {
    "project_code": "项目编号",
    "project_name": "项目名称",
    "project_type": "项目类型",
    "owner_name": "项目负责人",
    "department_name": "责任部门",
    "project_status": "项目状态",
    "planned_start": "计划开始",
    "planned_end": "计划结束",
}
NODE_MAPPING = {
    "project_code": "项目编号",
    "milestone_code": "节点编号",
    "milestone_name": "节点名称",
    "node_type_code": "节点类型",
    "planned_date": "计划完成日期",
    "department_name": "责任部门",
    "owner_name": "责任人",
    "source_ref": "日期依据",
}
STATUS_MAP = {
    "草稿": "draft",
    "待启动": "not_started",
    "进行中": "active",
    "已完成": "completed",
    "已归档": "archived",
}


def digest(value):
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()
    ).hexdigest()


def as_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value.strip() if isinstance(value, str) else value


def directory_id(db, cls, field, name):
    if not name:
        return None
    obj = db.scalar(select(cls).where(getattr(cls, field) == name))
    if not obj:
        fail(422, "DIRECTORY_VALUE_UNKNOWN", f"目录中未找到：{name}，请先配置演示目录")
    return obj.id


def normalize(db, target, raw, context):
    values = dict(raw)
    values["department_id"] = directory_id(
        db, Department, "dept_name", values.pop("department_name", None)
    )
    values["owner_id"] = directory_id(db, User, "display_name", values.pop("owner_name", None))
    if target == "projects":
        status = values.get("project_status") or "not_started"
        if status == "需关注":
            values["attention_note"] = "Excel 原始状态：需关注（生命周期未确认，暂存为草稿）"
            status = "draft"
        values["project_status"] = STATUS_MAP.get(status, status)
        return ProjectInput.model_validate(values)
    project = get(db, Project, context.get("project_id"))
    project_code = values.pop("project_code", None)
    if project_code and project_code != project.project_code:
        fail(422, "PROJECT_CODE_MISMATCH", "该行项目编号与选定项目不一致")
    plan = get(db, PlanVersion, context.get("plan_version_id"))
    if plan.project_id != project.id or plan.status != "draft":
        fail(409, "IMMUTABLE_PLAN", "节点导入必须选择本项目草稿计划")
    # Unconfirmed contract references are never silently promoted to contractual facts.
    values.update(
        plan_version_id=plan.id,
        date_kind="unconfirmed"
        if "未核实" in (values.get("source_ref") or "")
        else "internal_plan",
        change_reason="Excel导入草稿",
    )
    if values.get("tracking_only") in (None, "", "否"):
        values["tracking_only"] = False
    elif values.get("tracking_only") == "是":
        values["tracking_only"] = True
    return MilestoneInput.model_validate(values)


def duplicate(db, target, data, context):
    if target == "projects":
        return bool(db.scalar(select(Project.id).where(Project.project_code == data.project_code)))
    return bool(
        db.scalar(
            select(Milestone.id).where(
                Milestone.project_id == context["project_id"],
                Milestone.milestone_code == data.milestone_code,
            )
        )
    )


def preview(db, content, filename, target, mapping, sheet_name, header_row, context, request_id):
    if target not in ("projects", "milestones"):
        fail(422, "INVALID_TARGET", "仅支持 projects 或 milestones")
    if not filename.lower().endswith(".xlsx"):
        fail(400, "INVALID_FILE_TYPE", "仅支持 .xlsx 文件")
    if not content or len(content) > settings.max_upload_bytes:
        fail(413, "FILE_TOO_LARGE", "文件为空或超过 5MB")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 30 * 1024 * 1024:
                fail(413, "EXPANDED_FILE_TOO_LARGE", "Excel 解压后超过 30MB")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except (BadZipFile, ValueError, KeyError, OSError, ParseError, InvalidFileException):
        fail(400, "INVALID_WORKBOOK", "无法读取有效的 Excel 工作簿")
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        if sheet.max_row > settings.max_excel_rows + header_row:
            fail(413, "TOO_MANY_ROWS", "超出演示导入行数限制")
        defaults = PROJECT_MAPPING if target == "projects" else NODE_MAPPING
        mapping = mapping or defaults
        required = (
            {"project_code", "project_name"}
            if target == "projects"
            else {"milestone_code", "milestone_name", "node_type_code", "planned_date"}
        )
        allowed = set(defaults) | ({"tracking_only"} if target == "milestones" else set())
        if not required.issubset(mapping) or not set(mapping).issubset(allowed):
            fail(422, "INVALID_MAPPING", "字段映射缺少必填字段或含未知字段")
        headers = [
            str(c.value or "").strip()
            for c in next(sheet.iter_rows(min_row=header_row, max_row=header_row))
        ]
        if len(set(mapping.values())) != len(mapping) or any(
            headers.count(h) != 1 for h in mapping.values()
        ):
            fail(422, "INVALID_HEADERS", "映射列必须存在且唯一，请检查工作表和表头行号")
        batch = ImportBatch(
            target_entity=target,
            filename=filename[:255],
            file_sha256=hashlib.sha256(content).hexdigest(),
            mapping_json=mapping,
            context_json=context,
            preview_hash="pending",
            created_by=db.info.get("actor_id", "legacy-unassigned"),
        )
        db.add(batch)
        db.flush()
        seen = set()
        source_values = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            raw = {
                field: as_value(cells[headers.index(header)].value)
                for field, header in mapping.items()
            }
            if all(v is None or v == "" for v in raw.values()):
                continue
            errors, normalized, state = [], None, "valid"
            try:
                if any(
                    cells[headers.index(header)].data_type == "f" for header in mapping.values()
                ):
                    fail(422, "FORMULA_IN_INPUT", "业务输入列不接受公式，请提供原始值")
                data = normalize(db, target, raw, context)
                normalized = data.model_dump(mode="json")
                code = data.project_code if target == "projects" else data.milestone_code
                if code in seen or duplicate(db, target, data, context):
                    state = "duplicate"
                    errors = [
                        {
                            "field": "code",
                            "code": "DUPLICATE",
                            "message": "编号重复；不会覆盖已有记录",
                        }
                    ]
                seen.add(code)
            except ValidationError as exc:
                state = "error"
                errors = [
                    {
                        "field": ".".join(map(str, e["loc"])),
                        "code": "INVALID_VALUE",
                        "message": e["msg"],
                    }
                    for e in exc.errors()
                ]
            except HTTPException as exc:
                state = "error"
                errors = [{"field": "row", **exc.detail}]
            source_values.append(raw)
            db.add(
                ImportRow(
                    batch_id=batch.id,
                    sheet_name=sheet.title,
                    source_row_number=row_number,
                    raw_json=raw,
                    normalized_json=normalized,
                    validation_errors_json=errors,
                    row_status=state,
                )
            )
        if not source_values:
            fail(422, "EMPTY_WORKSHEET", "没有可导入的数据行")
        batch.preview_hash = digest(
            {
                "file": batch.file_sha256,
                "mapping": mapping,
                "context": context,
                "rows": source_values,
            }
        )
        audit(db, batch, "preview", request_id)
        return batch
    except KeyError:
        fail(422, "SHEET_NOT_FOUND", "工作表不存在，请检查工作表名称")
    finally:
        workbook.close()


def batch_detail(db, batch):
    result = record(batch)
    rows = db.scalars(
        select(ImportRow)
        .where(ImportRow.batch_id == batch.id)
        .order_by(ImportRow.source_row_number)
    ).all()
    result["rows"] = [record(r) for r in rows]
    result["valid_rows"] = sum(r.row_status == "valid" for r in rows)
    result["error_rows"] = sum(r.row_status in ("error", "duplicate") for r in rows)
    if batch.target_entity == "milestones" and not batch.result_json:
        project = get(db, Project, batch.context_json["project_id"])
        templates = list(db.scalars(select(ActionTemplate)))
        summary = {
            "automatic": 0,
            "conflict": 0,
            "missing": 0,
            "pending_approval": 0,
            "tracking_only": 0,
        }
        for row in rows:
            if row.row_status != "valid":
                continue
            item = SimpleNamespace(**row.normalized_json)
            candidates = [t for t in templates if wf.applicable(t, item, project)]
            state = (
                "tracking_only"
                if getattr(item, "tracking_only", False)
                else "conflict"
                if len(candidates) > 1
                else "automatic"
                if candidates
                else "pending_approval"
                if any(
                    t.status in ("draft", "pending_review") and wf.scope_applies(t, item, project)
                    for t in templates
                )
                else "missing"
            )
            summary[state] += 1
        result["matching_summary"] = summary
    result["attempts"] = [
        record(e)
        for e in db.scalars(
            select(AuditEvent)
            .where(AuditEvent.aggregate_id == batch.id, AuditEvent.action == "commit_rejected")
            .order_by(AuditEvent.occurred_at.desc())
        )
    ]
    return result


def commit(db, batch_id, data, key, request_id):
    batch = get(db, ImportBatch, batch_id, True)
    payload_hash = digest(data.model_dump())
    if batch.idempotency_key:
        if batch.idempotency_key == key and batch.commit_payload_hash == payload_hash:
            return batch
        fail(409, "IDEMPOTENCY_CONFLICT", "批次已提交；请勿用不同参数再次提交")
    if db.scalar(select(ImportBatch.id).where(ImportBatch.idempotency_key == key)):
        fail(409, "IDEMPOTENCY_CONFLICT", "该幂等键已用于另一批次")
    if batch.preview_hash != data.preview_hash:
        fail(409, "PREVIEW_CHANGED", "预览摘要不匹配，请重新预览")
    if data.include_valid_only:
        fail(409, "PARTIAL_IMPORT_NOT_CONFIRMED", "部分导入策略尚未确认；默认整批原子导入")
    rows = db.scalars(
        select(ImportRow)
        .where(ImportRow.batch_id == batch.id)
        .order_by(ImportRow.source_row_number)
    ).all()
    # Recheck all rows under project/batch locks, before any business insert.
    if batch.target_entity == "milestones":
        get(db, Project, batch.context_json["project_id"], True)
    valid = []
    for row in rows:
        if row.row_status != "valid":
            continue
        obj = normalize(db, batch.target_entity, row.raw_json, batch.context_json)
        if duplicate(db, batch.target_entity, obj, batch.context_json):
            fail(
                409,
                "IMPORT_DATA_CHANGED",
                f"第 {row.source_row_number} 行编号在预览后已存在，请重新预览",
            )
        valid.append((row, obj))
    if len(valid) != len(rows) and not data.include_valid_only:
        fail(
            409, "IMPORT_HAS_ERRORS", "预览含错误或重复行，请修正文件后重新上传；未写入任何业务数据"
        )
    if not valid:
        fail(409, "NO_VALID_ROWS", "没有合法数据可提交")
    matching_summary = {
        "automatic": 0,
        "conflict": 0,
        "missing": 0,
        "pending_approval": 0,
        "tracking_only": 0,
    }
    node_results, created_nodes = [], []
    for row, obj in valid:
        target = (
            create_project(db, obj, request_id, "excel_demo")
            if batch.target_entity == "projects"
            else create_milestone(
                db,
                batch.context_json["project_id"],
                obj,
                request_id,
                matching_source="automatic_excel_import",
            )
        )
        row.row_status = "inserted"
        row.target_record_id = target.id
        if batch.target_entity == "milestones":
            created_nodes.append(target)
    if created_nodes:
        project = get(db, Project, batch.context_json["project_id"])
        plan = get(db, PlanVersion, batch.context_json["plan_version_id"])
        items = {
            i.milestone_id: i
            for i in db.scalars(
                select(PlanItem).where(
                    PlanItem.plan_version_id == plan.id,
                    PlanItem.milestone_id.in_([n.id for n in created_nodes]),
                )
            )
        }
        matching_data = action_matching.catalog(db, [i.id for i in items.values()])
        configuration_data = node_configuration.catalog(db, project.id)
        for target in created_nodes:
            item = items[target.id]
            state = action_matching.describe(
                project,
                plan,
                item,
                target,
                db.info["principal"],
                configuration_data["roles"],
                matching_data,
            )
            configuration = node_configuration.describe(
                db, project, plan, item, target, state, configuration_data
            )
            matching_summary[state["status"]] += 1
            node_results.append(
                {
                    "milestone_id": target.id,
                    "milestone_code": target.milestone_code,
                    "status": state["status"],
                    "label": state["label"],
                    "action_count": state["action_count"],
                    "configuration_status": configuration["status"],
                    "configuration_label": configuration["label"],
                    "issues": configuration["issues"],
                }
            )
    batch.idempotency_key = key
    batch.commit_payload_hash = payload_hash
    batch.committed_at = now()
    batch.status = "committed" if len(valid) == len(rows) else "partial"
    batch.result_json = {"inserted": len(valid), "skipped": len(rows) - len(valid)}
    if batch.target_entity == "milestones":
        batch.result_json.update(matching_summary=matching_summary, node_results=node_results)
    audit(db, batch, "commit", request_id)
    return batch
